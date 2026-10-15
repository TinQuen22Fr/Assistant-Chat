import io
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader

MAX_FILE_MB = 20
MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024
MAX_TEXT_CHARS = 400_000
MAX_IMAGE_SIDE = 2048
ACCEPTED_LABEL = "PDF, DOCX, TXT, MD, PNG, JPG, WEBP"

IMAGE_TYPES = {
    "png": ("image/png", "PNG"),
    "jpg": ("image/jpeg", "JPEG"),
    "jpeg": ("image/jpeg", "JPEG"),
    "webp": ("image/webp", "WEBP"),
}
DOC_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "md": "text/markdown",
}


class AttachmentError(Exception):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.message = message
        self.status = status


@dataclass
class Attachment:
    name: str
    kind: str  # "image" | "pdf" | "text"
    mime: str
    data: bytes = b""
    text: str = ""


@dataclass
class Turn:
    role: str  # "user" | "model"
    text: str
    attachments: list[Attachment] = field(default_factory=list)
    ts: object = None


def _mb(size: int) -> str:
    return f"{size / 1024 / 1024:.1f}".replace(".", ",")


def ext_of(name: str) -> str:
    return Path(name).suffix.lower().lstrip(".")


def normalize_image(name: str, data: bytes) -> bytes:
    fmt = IMAGE_TYPES[ext_of(name)][1]
    try:
        img = Image.open(io.BytesIO(data))
        img.seek(0)
        img.load()
    except Exception:
        raise AttachmentError(f"L'image « {name} » est illisible ou corrompue.")
    if img.mode == "P":
        img = img.convert("RGBA")
    if fmt == "JPEG" and img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    out = io.BytesIO()
    img.save(out, format=fmt)
    return out.getvalue()


def _pdf_text(name: str, data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    except Exception:
        raise AttachmentError(f"Le PDF « {name} » est illisible ou corrompu.")


def _docx_text(name: str, data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise AttachmentError(f"Le document Word « {name} » est illisible ou corrompu.")
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def parse_upload(name: str, data: bytes) -> Attachment:
    """Validate an upload and extract cacheable text."""
    ext = ext_of(name)
    if ext not in IMAGE_TYPES and ext not in DOC_TYPES:
        raise AttachmentError(
            f"Format non pris en charge pour « {name} ». Formats acceptés : {ACCEPTED_LABEL}.", 415
        )
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentError(
            f"Le fichier « {name} » dépasse la limite autorisée de {MAX_FILE_MB} Mo ({_mb(len(data))} Mo).", 413
        )
    if not data:
        raise AttachmentError(f"Le fichier « {name} » est vide.")

    if ext in IMAGE_TYPES:
        normalize_image(name, data)
        return Attachment(name, "image", IMAGE_TYPES[ext][0])
    if ext == "pdf":
        return Attachment(name, "pdf", DOC_TYPES["pdf"], text=_pdf_text(name, data)[:MAX_TEXT_CHARS])
    text = _docx_text(name, data) if ext == "docx" else _decode_text(data)
    return Attachment(name, "text", DOC_TYPES[ext], text=text[:MAX_TEXT_CHARS])


def load_attachment(doc: dict, raw: bytes | None) -> Attachment:
    """Rebuild a model-ready attachment from a stored file document."""
    if doc["kind"] == "image":
        return Attachment(doc["name"], "image", doc["mime"], data=normalize_image(doc["name"], raw))
    if doc["kind"] == "pdf":
        return Attachment(doc["name"], "pdf", doc["mime"], data=raw)
    return Attachment(doc["name"], "text", doc["mime"], text=doc.get("text", ""))


def turn_text(turn: Turn) -> str:
    if turn.role != "user":
        return turn.text
    blocks = [
        f"[Document joint : {a.name}]\n<<<\n{a.text}\n>>>"
        for a in turn.attachments
        if a.kind == "text"
    ]
    names = [a.name for a in turn.attachments if a.kind != "text"]
    if names:
        blocks.append("[Fichiers joints : " + ", ".join(names) + "]")
    blocks.append(turn.text or "Analyse les fichiers joints et résume leur contenu.")
    return "\n\n".join(blocks)
